"""File upload endpoint — extract text from files and ingest into RAG pipeline.

Supported formats: PDF, DOCX, PPTX, HTML, TXT, Markdown.
Server-side parsing avoids fragile client-side extraction.
"""

from __future__ import annotations

import asyncio
import io
import logging
import uuid as uuid_mod
from pathlib import Path

from fastapi import APIRouter, Header, HTTPException, Request, UploadFile
from haystack import Document
from pydantic import BaseModel

from agentic_ra.agent.tools import _sanitize_user_id
from agentic_ra.pipelines.factory import RAGPipeline

logger = logging.getLogger(__name__)

router = APIRouter(prefix="/api/v1", tags=["upload"])

MAX_FILE_SIZE = 50 * 1024 * 1024  # 50 MB

ALLOWED_EXTENSIONS = frozenset({".pdf", ".docx", ".pptx", ".html", ".htm", ".txt", ".md"})


class UploadResponse(BaseModel):
    """File upload + ingestion result."""

    ingested_count: int
    chunk_count: int
    parsed_document_id: uuid_mod.UUID
    cognee_ingested: bool
    filename: str
    file_type: str


def _get_pipeline(request: Request) -> RAGPipeline:
    return request.app.state.pipeline  # type: ignore[no-any-return]


def _extract_text_pdf(content: bytes) -> str:
    """Extract text from PDF using PyMuPDF."""
    try:
        import pymupdf
    except ImportError as exc:
        raise HTTPException(
            status_code=501,
            detail="PDF parsing unavailable — pymupdf not installed",
        ) from exc

    doc = pymupdf.open(stream=content, filetype="pdf")
    pages = [page.get_text() for page in doc]
    doc.close()
    return "\n\n".join(pages)


def _extract_text_docx(content: bytes) -> str:
    """Extract text from DOCX using python-docx."""
    try:
        from docx import Document as DocxDocument
    except ImportError as exc:
        raise HTTPException(
            status_code=501,
            detail="DOCX parsing unavailable — python-docx not installed",
        ) from exc

    doc = DocxDocument(io.BytesIO(content))
    return "\n\n".join(p.text for p in doc.paragraphs if p.text.strip())


def _extract_text_pptx(content: bytes) -> str:
    """Extract text from PPTX using python-pptx."""
    try:
        from pptx import Presentation
    except ImportError as exc:
        raise HTTPException(
            status_code=501,
            detail="PPTX parsing unavailable — python-pptx not installed",
        ) from exc

    prs = Presentation(io.BytesIO(content))
    parts: list[str] = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                parts.append(shape.text_frame.text)
    return "\n\n".join(parts)


def _extract_text_html(content: bytes) -> str:
    """Extract text from HTML, stripping tags."""
    try:
        from bs4 import BeautifulSoup
    except ImportError:
        return content.decode("utf-8", errors="replace")

    soup = BeautifulSoup(content.decode("utf-8", errors="replace"), "html.parser")
    for tag in soup(["script", "style", "nav", "footer", "header"]):
        tag.decompose()
    return soup.get_text(separator="\n", strip=True)


_EXTRACTORS: dict[str, object] = {
    ".pdf": _extract_text_pdf,
    ".docx": _extract_text_docx,
    ".pptx": _extract_text_pptx,
    ".html": _extract_text_html,
    ".htm": _extract_text_html,
}


def _extract_text(filename: str, content: bytes) -> str:
    """Route to the correct text extractor based on file extension."""
    ext = Path(filename).suffix.lower()
    extractor = _EXTRACTORS.get(ext)
    if extractor is not None:
        return extractor(content)  # type: ignore[operator]
    # Plain text / markdown — decode directly
    return content.decode("utf-8", errors="replace")


@router.post("/upload", response_model=UploadResponse)
async def upload_endpoint(
    request: Request,
    file: UploadFile,
    x_user_id: str = Header(default="anonymous", alias="X-User-Id"),
) -> UploadResponse:
    """Upload a file, extract text, and ingest into the RAG pipeline."""
    filename = file.filename or "unknown"
    ext = Path(filename).suffix.lower()

    if ext not in ALLOWED_EXTENSIONS:
        raise HTTPException(
            status_code=415,
            detail=(
                f"Unsupported file type '{ext}'. "
                f"Allowed: {', '.join(sorted(ALLOWED_EXTENSIONS))}"
            ),
        )

    content = await file.read()

    if len(content) > MAX_FILE_SIZE:
        raise HTTPException(
            status_code=413,
            detail=f"File exceeds {MAX_FILE_SIZE // (1024 * 1024)} MB limit",
        )
    if not content:
        raise HTTPException(status_code=422, detail="Empty file")

    text = _extract_text(filename, content)
    if not text.strip():
        raise HTTPException(status_code=422, detail="No text content extracted from file")

    pipeline = _get_pipeline(request)
    dataset = f"user_{_sanitize_user_id(x_user_id)}"

    docs = [Document(content=text, meta={"title": filename, "source": "upload"})]

    chunk_result = await asyncio.to_thread(pipeline.embedding_pipeline.chunk_text, docs)

    docling_json = (
        chunk_result.docling_doc.export_to_dict()
        if chunk_result.docling_doc is not None
        else {}
    )
    parsed_doc = await pipeline.document_store.save(
        source_path=f"upload/{filename}",
        docling_json=docling_json,
        chunk_count=chunk_result.chunk_count,
        metadata={"filename": filename, "file_type": ext},
        user_id=x_user_id,
    )

    cognee_ok = False
    if pipeline.cognee_store is not None:
        try:
            await pipeline.cognee_store.add_and_cognify(text=text, dataset=dataset)
            cognee_ok = True
            logger.info("Cognee ingestion complete for upload: %s", filename)
        except Exception:
            logger.exception("Cognee add+cognify failed for upload (non-fatal)")

    return UploadResponse(
        ingested_count=1,
        chunk_count=chunk_result.chunk_count,
        parsed_document_id=parsed_doc.id,
        cognee_ingested=cognee_ok,
        filename=filename,
        file_type=ext,
    )
